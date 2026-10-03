"""Add summary, scenario, sensitivity, Monte Carlo, risk and KPI sheets to the model.
Static tables are snapshots from twin.py/analysis.py/stress.py (base inputs, 03.10.2026)."""
import json
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.utils import get_column_letter as CL
BLK=Font(name='Arial',size=10); BOLD=Font(name='Arial',size=10,bold=True); TITLE=Font(name='Arial',size=14,bold=True)
H2=Font(name='Arial',size=11,bold=True,color='2C6A55'); NOTE=Font(name='Arial',size=9,italic=True,color='5A6762'); GRN=Font(name='Arial',size=10,color='008000')
HDR=PatternFill('solid',fgColor='E2EEE8'); TOT=PatternFill('solid',fgColor='F1F5F2'); RED=PatternFill('solid',fgColor='F6E7DE'); YEL=PatternFill('solid',fgColor='FFF4C2')
USD='$#,##0;($#,##0);"-"'; PCT='0.0%;(0.0%);"-"'; NUM='#,##0;(#,##0);"-"'
A=json.load(open('analysis.json')); L=json.load(open('layout.json')); rows=L['rows']; YR=L['YROW']
wb=load_workbook('Aldera_FinModel.xlsx')
SNAP='Снимок расчёта от 03.10.2026 по базовым вводным. После изменения допущений пересчитать: python3 analysis.py && python3 stress.py && python3 build_analysis_sheets.py'
def hdr(ws,r,items):
    for i,h in enumerate(items): c=ws.cell(r,1+i,h); c.font=BOLD; c.fill=HDR
def put(ws,r,vals,fmts=None,font=BLK,fill=None):
    for i,v in enumerate(vals):
        c=ws.cell(r,1+i,v); c.font=font
        if fmts and i<len(fmts) and fmts[i]: c.number_format=fmts[i]
        if fill: c.fill=fill
for name in ('Сводка','Сценарии','Чувствительность','Монте-Карло','Риски','KPI по кварталам'):
    if name in wb.sheetnames: del wb[name]

# ---- Сводка (live links) ----
S=wb.create_sheet('Сводка',0)
S['A1']='The Aldera · финансовая модель · сводка'; S['A1'].font=TITLE
S['A2']='518 GROUP LLC. 4 продукта Vox Nutrition, каналы TikTok Shop, Shopify, Amazon. Старт продаж 20.11.2026. Вложение двумя траншами. Платная реклама только Spark Ads на выстрелившие ролики, с февраля 2027.'; S['A2'].font=NOTE
S['A3']="Активный сценарий (1 конс., 2 база, 3 опт.):"; S['A3'].font=BOLD; S['C3']="='Допущения'!B3"; S['C3'].font=GRN
hdr(S,5,['Показатель','Год 1','Год 2','Год 3'])
Y="'Итоги по годам'"
items=[('Выручка','rev',USD),('Маржинальная прибыль','contrib',USD),('Маржинальность','cmp',PCT),('Постоянные расходы','opex',USD),('EBITDA','ebitda',USD),('Рентабельность по EBITDA','em',PCT),('Чистая прибыль после налога (оценка)','net',USD),('Заказов всего','ord',NUM),('Подписчиков на конец года','subs_end',NUM),('Криейторов на конец года','cr_end',NUM),('Деньги на счёте на конец года','cash_end',USD)]
r=6
for lab,k,fmt in items:
    S.cell(r,1,lab).font=BOLD if k in('rev','ebitda') else BLK
    for j in range(3):
        c=S.cell(r,2+j,f"={Y}!{CL(2+j)}{YR[k]}"); c.number_format=fmt; c.font=GRN
        if k in('rev','ebitda'): c.fill=TOT
    r+=1
r+=1
S.cell(r,1,'Минимум денег на счёте').font=BOLD; c=S.cell(r,2,f"={Y}!E{YR['mincash']}"); c.number_format=USD; c.font=GRN; mc_row=r; r+=1
S.cell(r,1,'Месяц выхода на темп $1 млн EBITDA в год (1 = ноябрь 2026)').font=BOLD; c=S.cell(r,2,f"={Y}!E{YR['m1m']}"); c.font=GRN; r+=2
S.cell(r,1,'Главные выводы (базовый сценарий, снимок 03.10.2026)').font=H2; r+=1
b=A['scen_base']; mc=A['mc']
fm=lambda v: f"{v:,.0f}".replace(',',' ')
import math
T2=int(math.ceil((A['mc']['need_p90']+5000)/1000.0)*1000)
MN=['','ноябрь 2026','декабрь 2026','январь 2027','февраль 2027','март 2027','апрель 2027','май 2027','июнь 2027','июль 2027','август 2027','сентябрь 2027','октябрь 2027']
facts=[f"Вложения окупаются EBITDA на {b['payback']}-м месяце ({MN[b['payback']]}), темп $1 млн EBITDA в год на {b['m1m']}-м месяце ({MN[b['m1m']]}).",
 f"Узкое место {MN[b['mincash_m']]}: при запасе товара 1,3 месяца денег не хватает на ${fm(-b['mincash'])}. Решение: второй транш ${fm(T2)} вместо $5 000 (покрывает 90% исходов Монте-Карло)"+(' или запас 1,0 месяца в декабре.' if A['funding'][1][1]>=0 else f'. Запас 1,0 месяца не спасает (нехватка ${fm(-A["funding"][1][1])}): нужны деньги владельца, отсрочка Vox или кредит под товар.'),
 f"Монте-Карло, {fm(A['mc_n'])} прогонов: EBITDA года 1 от ${fm(mc['ebitda1']['p10']/1e3)} тыс. (P10) до ${fm(mc['ebitda1']['p90']/1e3)} тыс. (P90), медиана ${fm(mc['ebitda1']['p50']/1e3)} тыс.",
 f"Вероятность EBITDA года 2 не меньше $1 млн: {mc['p_y2_1m']*100:.1f}%. Вероятность, что при транше $5 000 денег не хватит: {mc['p_cash_neg']*100:.0f}%, нехватка до ${fm(mc['need_p90'])} (P90).".replace('.',',',1) if False else f"Вероятность EBITDA года 2 не меньше $1 млн: {('больше 99' if mc['p_y2_1m']>0.99 else str(round(mc['p_y2_1m']*100,1)).replace('.',','))}%. Вероятность, что при транше $5 000 денег не хватит: {mc['p_cash_neg']*100:.0f}%, нехватка до ${fm(mc['need_p90'])} (P90).",
 "Главные рычаги прибыли: конверсия просмотра в заказ, просмотры на ролик, скорость набора криейторов, переход в подписку.",
 f"Без Amazon EBITDA года 1 остаётся в плюсе уже при конверсии {str(round(A['be_conv_y1_noamz']*100,4)).replace('.',',')}%; $1 млн EBITDA во 2-м году без Amazon при конверсии от {str(round(A['be_conv_y2_1m_noamz']*100,3)).replace('.',',')}%."]
for f in facts: S.cell(r,1,'• '+f).font=BLK; r+=1
r+=1
# chart: monthly revenue & EBITDA
M=wb['Модель']
ch=BarChart(); ch.type='col'; ch.title='Выручка и EBITDA по месяцам, $'; ch.height=8; ch.width=26
data=Reference(M,min_col=4,max_col=39,min_row=rows['rev'],max_row=rows['rev']); ch.add_data(data,from_rows=True,titles_from_data=False)
data2=Reference(M,min_col=4,max_col=39,min_row=rows['ebitda'],max_row=rows['ebitda']); ch.add_data(data2,from_rows=True,titles_from_data=False)
ch.series[0].tx=None; ch.series[1].tx=None
from openpyxl.chart.series import SeriesLabel
ch.series[0].tx=SeriesLabel(v='Выручка'); ch.series[1].tx=SeriesLabel(v='EBITDA')
ch.set_categories(Reference(M,min_col=4,max_col=39,min_row=3,max_row=3))
S.add_chart(ch,f'A{r}')
S.column_dimensions['A'].width=70
for cc in 'BCD': S.column_dimensions[cc].width=16

# ---- KPI по кварталам (live) ----
K=wb.create_sheet('KPI по кварталам',2)
K['A1']='KPI по кварталам: цели для ежемесячного контроля'; K['A1'].font=TITLE
K['A2']='Квартал 1 = ноябрь 2026 - январь 2027. Живые формулы из листа «Модель».'; K['A2'].font=NOTE
QN=['Q1 ноя26-янв27','Q2 фев-апр27','Q3 май-июл27','Q4 авг-окт27','Q5 ноя27-янв28','Q6 фев-апр28','Q7 май-июл28','Q8 авг-окт28','Q9','Q10','Q11','Q12']
hdr(K,4,['Показатель']+QN)
KL=[('Криейторов на конец квартала','cr','end',NUM),('Роликов вышло','vid','sum',NUM),('Просмотров, млн','views','summ',"0.0"),('Заказы TikTok','tt','sum',NUM),('Заказы Shopify первые','sd','sum',NUM),('Заказы подписки','subord','sum',NUM),('Подписчиков на конец','subs','end',NUM),('Заказы Amazon','az','sum',NUM),
    ('Выручка','rev','sum',USD),('Маржинальная прибыль','contrib','sum',USD),('EBITDA','ebitda','sum',USD),('Закупка у Vox','o_buy','sum',USD),('Деньги на конец квартала','cash','end',USD)]
r=5
for lab,k,how,fmt in KL:
    K.cell(r,1,lab).font=BOLD if k in('rev','ebitda') else BLK
    for q in range(12):
        c1=CL(4+3*q); c3=CL(4+3*q+2); rr=rows[k]
        f={'sum':f"=SUM('Модель'!{c1}{rr}:{c3}{rr})",'end':f"='Модель'!{c3}{rr}",'summ':f"=SUM('Модель'!{c1}{rr}:{c3}{rr})/1000000"}[how]
        c=K.cell(r,2+q,f); c.number_format=fmt; c.font=GRN
    r+=1
K.column_dimensions['A'].width=34
for q in range(12): K.column_dimensions[CL(2+q)].width=14
K.freeze_panes='B5'

# ---- Сценарии ----
C=wb.create_sheet('Сценарии')
C['A1']='Три сценария'; C['A1'].font=TITLE; C['A2']=SNAP; C['A2'].font=NOTE
hdr(C,4,['Показатель','Консервативный','Базовый','Оптимистичный'])
sc=[A['scen_cons'],A['scen_base'],A['scen_opt']]
lines=[('Выручка год 1','rev1',USD),('EBITDA год 1','ebitda1',USD),('Выручка год 2','rev2',USD),('EBITDA год 2','ebitda2',USD),('Выручка год 3','rev3',USD),('EBITDA год 3','ebitda3',USD),
       ('Чистая прибыль 3 года (сумма)',None,USD),('Подписчиков через 12 мес.','subs1',NUM),('Подписчиков через 24 мес.','subs2',NUM),('Криейторов через 24 мес.','cr2',NUM),
       ('Деньги на счёте через 12 мес.','cash1',USD),('Минимум денег на счёте','mincash',USD),('Месяц темпа $1 млн EBITDA в год','m1m',None),('Окупаемость $25 000, месяц','payback',None)]
r=5
for lab,k,fmt in lines:
    vals=[lab]+[ (s['net1']+s['net2']+s['net3']) if k is None else (s[k] if s[k] is not None else 'нет за 36 мес.') for s in sc]
    put(C,r,vals,[None,fmt,fmt,fmt]); r+=1
r+=1; C.cell(r,1,'Что отличает сценарии').font=H2; r+=1
hdr(C,r,['Параметр','Консервативный','Базовый','Оптимистичный']); r+=1
for lab,a,b2,c3 in [('Просмотров на ролик','5 000','15 000','30 000'),('Конверсия','0,03%','0,05%','0,07%'),('Доля вышедших роликов','60%','100%','100%'),('Переход в подписку','12%','15%','20%'),('Отток подписки','12%','10%','8%'),('Новых криейторов в мес.','5-6','10-12','15-18'),('Amazon, продаж в день на продукт','13','30','45'),('Spark Ads, выручка на $1 рекламы','2','3','4'),('Цена Vox','+5%','прайс','прайс')]:
    put(C,r,[lab,a,b2,c3]); r+=1
r+=1; C.cell(r,1,'Деньги на старте: варианты финансирования (базовый сценарий)').font=H2; r+=1
hdr(C,r,['Вариант','Минимум денег','Месяц минимума','EBITDA год 1']); r+=1
for lab,mn,mm,e1 in A['funding']:
    put(C,r,[lab,mn,mm,e1],[None,USD,None,USD],fill=RED if mn<0 else None); r+=1
C.column_dimensions['A'].width=46
for cc in 'BCD': C.column_dimensions[cc].width=18

# ---- Чувствительность ----
T=wb.create_sheet('Чувствительность')
T['A1']='Чувствительность и стресс-тесты'; T['A1'].font=TITLE; T['A2']=SNAP; T['A2'].font=NOTE
T['A4']='Торнадо: каждый параметр отдельно, остальное по базе'; T['A4'].font=H2
hdr(T,5,['Параметр','Плохое значение','Хорошее значение','EBITDA год 1, плохо','EBITDA год 1, хорошо','EBITDA год 2, плохо','EBITDA год 2, хорошо','Размах год 2'])
r=6
for lab,lo,hi,e2lo,e2hi,e1lo,e1hi in A['tornado']:
    f=PCT if (isinstance(lo,float) and lo<1 and lab not in('Цена Vox',)) else ('0.00x' if lab=='Цена Vox' else NUM)
    if lab=='Конверсия просмотра в заказ': f='0.00%'
    put(T,r,[lab,lo,hi,e1lo,e1hi,e2lo,e2hi,abs(e2hi-e2lo)],[None,f,f,USD,USD,USD,USD,USD]); r+=1
r+=1; T.cell(r,1,'Стресс-тесты').font=H2; r+=1
hdr(T,r,['Событие','EBITDA год 1','EBITDA год 2','EBITDA год 3','Минимум денег','Темп $1 млн, месяц']); r+=1
for s in A['stress']:
    put(T,r,[s['lab'],s['e1'],s['e2'],s['e3'],s['mincash'],s['m1m'] or 'нет'],[None,USD,USD,USD,USD,None],fill=TOT if s['lab']=='Базовый сценарий' else None); r+=1
r+=1; T.cell(r,1,'Точки безубыточности по конверсии (без Amazon, остальное по базе)').font=H2; r+=1
for lab,v in (('EBITDA года 1 больше нуля при конверсии от',A['be_conv_y1_noamz']),('EBITDA года 2 от $1 млн при конверсии от',A['be_conv_y2_1m_noamz']),('EBITDA года 1 от $1 млн при конверсии от',A['be_conv_y1_1m_noamz'])):
    put(T,r,[lab,v],[None,'0.000%']); r+=1
T.cell(r,1,'База: 0,05%. Запас прочности по конверсии большой: бизнес остаётся прибыльным, даже если конверсия в 10 раз ниже базы.').font=NOTE
T.column_dimensions['A'].width=62
for cc in 'BCDEFGH': T.column_dimensions[cc].width=17

# ---- Монте-Карло ----
MCW=wb.create_sheet('Монте-Карло')
MCW['A1']=f"Монте-Карло: {A['mc_n']} случайных прогонов модели"; MCW['A1'].font=TITLE; MCW['A2']=SNAP+' Сид 42.'; MCW['A2'].font=NOTE
MCW['A4']='Распределения вводных (треугольные: минимум, наиболее вероятное, максимум)'; MCW['A4'].font=H2
hdr(MCW,5,['Параметр','Минимум','Наиболее вероятное','Максимум'])
DN={'views':'Просмотров на ролик','conv':'Конверсия','deliv':'Доля вышедших роликов','org':'Органика','direct':'Прямые заказы Shopify','subconv':'Переход в подписку','churn':'Отток подписки','add1':'Новых криейторов мес. 3-12','add2':'Новых криейторов с 13-го','azpeak':'Amazon, продаж в день на продукт','azgrow':'Amazon, прирост в день в мес.','ppc':'Amazon, доля рекламы','sroas':'Spark Ads, ROAS','shock':'Множитель цены Vox'}
r=6
for k,(a,b2,c3) in A['mc_dist'].items():
    f='0.00%' if k=='conv' else (PCT if k in('deliv','org','direct','subconv','churn','ppc') else ('0.00x' if k in('shock','sroas') else '#,##0.0'))
    put(MCW,r,[DN[k],a,b2,c3],[None,f,f,f]); r+=1
r+=1; MCW.cell(r,1,'Результаты').font=H2; r+=1
hdr(MCW,r,['Показатель','P10 (плохой из 10)','P50 (медиана)','P90 (хороший из 10)','Среднее']); r+=1
for lab,k in (('Выручка год 1','rev1'),('EBITDA год 1','ebitda1'),('Выручка год 2','rev2'),('EBITDA год 2','ebitda2'),('Выручка год 3','rev3'),('EBITDA год 3','ebitda3'),('Подписчиков через 24 мес.','subs2'),('Минимум денег на счёте','mincash')):
    d=mc[k]; put(MCW,r,[lab,d['p10'],d['p50'],d['p90'],d['mean']],[None]+[NUM if k=='subs2' else USD]*4); r+=1
r+=1; MCW.cell(r,1,'Вероятности').font=H2; r+=1
for lab,v,f in (('EBITDA года 1 в минусе',mc['p_y1_loss'],PCT),('EBITDA года 1 от $1 млн',mc['p_y1_1m'],PCT),('EBITDA года 2 от $1 млн',mc['p_y2_1m'],PCT),
                ('Темп $1 млн EBITDA в год наступает в первые 24 месяца',mc['p_m1m_within24'],PCT),('Денег на старте не хватает без доплаты',mc['p_cash_neg'],PCT),('Не хватает больше $10 000',mc['p_cash_neg10k'],PCT),
                ('Нужная доплата с вероятностью 90%',mc['need_p90'],USD),('Месяц темпа $1 млн: медиана',mc['m1m_p50'],None),('Месяц темпа $1 млн: P10-P90',f"{mc['m1m_p10']}-{mc['m1m_p90']}",None)):
    put(MCW,r,[lab,v],[None,f]); r+=1
r+=1; MCW.cell(r,1,'Распределение EBITDA года 2').font=H2; r+=1
hdr(MCW,r,['Диапазон','Доля прогонов']); r+=1; h0=r
for lab,v in mc['hist_y2']: put(MCW,r,[lab,v],[None,PCT]); r+=1
bc=BarChart(); bc.title='EBITDA года 2: доля прогонов'; bc.height=7; bc.width=16
bc.add_data(Reference(MCW,min_col=2,min_row=h0,max_row=r-1),titles_from_data=False); bc.set_categories(Reference(MCW,min_col=1,min_row=h0,max_row=r-1)); bc.legend=None
MCW.add_chart(bc,f'D{h0-1}')
MCW.column_dimensions['A'].width=56
for cc in 'BCDE': MCW.column_dimensions[cc].width=18

# ---- Риски ----
RK=wb.create_sheet('Риски')
RK['A1']='Реестр рисков'; RK['A1'].font=TITLE
RK['A2']='Вероятность это оценка на 03.10.2026. Ущерб: потеря EBITDA года 1 из стресс-тестов и торнадо. Ожидаемый ущерб = вероятность × ущерб.'; RK['A2'].font=NOTE
hdr(RK,4,['Риск','Вероятность','Ущерб, EBITDA год 1','Ожидаемый ущерб','Ранний сигнал (KPI)','Что делаем','Кто'])
st={s['lab']:s for s in A['stress']}; base1=A['scen_base']['ebitda1']
tor={t[0]:t for t in A['tornado']}
RISKS=[('Криейторы не выпускают 8 роликов',0.50,base1-tor['Доля вышедших роликов'][5],'меньше 2 роликов в неделю на криейтора','еженедельный учёт, замена неактивных, бонус лучшим','Полина'),
 ('Конверсия ниже базы (0,03%)',0.30,base1-tor['Конверсия просмотра в заказ'][5],'меньше 0,4 заказа на 1 000 просмотров в первые 2 недели','тест хуков, Flip the jar, стартовая скидка','владелец, Полина'),
 ('Подписка слабее плана',0.25,base1-st['Подписка не работает: переход 5%, отток 15%']['e1'],'переход в подписку меньше 10% за первые 60 дней','письма день 0, 7, 25, скидка на первую подписку, вкладыш','Claude, владелец'),
 ('Amazon не даёт допуск к добавкам',0.20,base1-st['Amazon не дал допуск к категории']['e1'],'нет допуска к 15.02.2027','COA из лаборатории ISO 17025 в январе, Brand Registry по заявке USPTO','владелец'),
 ('Криейторов не прибавляется',0.20,base1-st['Криейторов не прибавляется: 50 весь срок']['e1'],'меньше 8 новых в месяц','аутрич 100 в неделю, менеджер криейторов с 5-го месяца','Полина'),
 ('TikTok Shop блокирует магазин на 2 месяца',0.10,base1-st['TikTok Shop заблокирован на 2 месяца (март-апрель 2027)']['e1'],'страйки за заявления, жалобы','этикетка завода дословно, без заявлений о болезни, Shopify и Amazon как опора','Claude'),
 ('Цены Vox растут на 15%',0.15,base1-st['Цены Vox +15%']['e1'],'новый прайс Vox','годовой объём под цену, SMP вторым источником','владелец'),
 ('Vox срывает отгрузку на месяц',0.15,base1-st['Vox не отгрузил товар: месяц без продаж (февраль 2027)']['e1'],'срок перезаказа больше 15 рабочих дней','запас 1,3 месяца, подтверждение мощности в мае 2027, SMP в резерве','Claude'),
 ('Денег на товар не хватает в декабре',0.85,22615*0.38,'деньги на счёте меньше $3 000',f'второй транш ${fm(T2)} или перезаказ дважды в месяц','владелец'),
 ('Задержка выплат TikTok новому продавцу',0.30,0,'выплата позже 14 дней после доставки','резерв в кассе, Shopify как быстрый канал денег','владелец'),
 ('Претензия по качеству или отзыв партии',0.05,50000,'жалобы на вкус, комки, реакции','COA каждой партии, страховка ответственности, этикетка завода','владелец'),
 ('Нагрузка на владельца и Полину',0.40,0,'ответ клиенту дольше 24 часов','найм с 5-6-го месяца по плану команды','владелец')]
r=5
for lab,p,dmg,kpi,act,who in RISKS:
    put(RK,r,[lab,p,dmg,f'=B{r}*C{r}',kpi,act,who],[None,PCT,USD,USD]); 
    if p*dmg>50000: 
        for cc in range(1,8): RK.cell(r,cc).fill=RED
    r+=1
RK.cell(r,1,'Итого ожидаемый ущерб').font=BOLD; c=RK.cell(r,4,f'=SUM(D5:D{r-1})'); c.number_format=USD; c.font=BOLD
RK.column_dimensions['A'].width=40; RK.column_dimensions['B'].width=13; RK.column_dimensions['C'].width=18; RK.column_dimensions['D'].width=17
RK.column_dimensions['E'].width=44; RK.column_dimensions['F'].width=60; RK.column_dimensions['G'].width=16

order=['Сводка','Итоги по годам','KPI по кварталам','Юнит-экономика','Сценарии','Чувствительность','Монте-Карло','Риски','Допущения','Модель']
wb._sheets=[wb[n] for n in order]
wb.active=0
wb.save('Aldera_FinModel.xlsx'); print('ok')
