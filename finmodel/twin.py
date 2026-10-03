"""Python twin of Aldera_FinModel.xlsx. Same logic as the sheet formulas; used for scenarios,
sensitivity and Monte Carlo. Verified against the recalculated workbook (see verify())."""
import math, bisect
TIERS=[150,300,500,1000,2500,5000,7500,10000]
SK=['COL','PRO','MG','CRE']
PROD={'COL':dict(ret=39.95,sh=4.50,fba=6.55,mix=.35,fb=1000,p=[8.52,8.37,8.12,7.97,7.82,7.67,7.42,7.12]),
      'PRO':dict(ret=34.95,sh=3.50,fba=5.40,mix=.15,fb=400,p=[8.34,7.99,7.64,7.44,7.24,6.94,6.74,6.59]),
      'MG': dict(ret=29.95,sh=3.50,fba=5.40,mix=.25,fb=1000,p=[5.51,5.16,4.81,4.61,4.41,4.11,3.91,3.76]),
      'CRE':dict(ret=29.95,sh=5.50,fba=6.55,mix=.25,fb=200,p=[6.90,6.75,6.50,6.35,6.20,6.05,5.80,5.50])}
OFF=[(59.95,5.50,.35,(1,0,1,0)),(64.95,5.50,.10,(1,1,0,0)),(114.95,8.50,.05,(1,1,1,1)),(39.95,4.50,.20,(1,0,0,0)),(34.95,3.50,.12,(0,1,0,0)),(29.95,3.50,.10,(0,0,1,0)),(29.95,5.50,.08,(0,0,0,1))]
SUBS=[(54.95,5.50,.60,(1,0,1,0)),(59.95,5.50,.25,(1,1,0,0)),(99.95,7.50,.15,(1,1,1,.5))]
SEAS=[1.10,1,1,1,1,1,1,1,1,1,1.15,1.15]
TEAM=[(4000,5),(3000,6),(5000,10),(4000,13),(3000,16),(2000,18)]
SCEN={ # cons, base, opt
 'v1':(10700,12700,17500),'v2':(6300,7800,9400),'v3':(450,560,2000),'conv':(.0003,.0005,.0007),'deliv':(.6,1,1),'org':(.1,.2,.3),'direct':(.05,.1,.15),
 'subconv':(.12,.15,.2),'churn':(.12,.1,.08),'add1':(5,10,15),'add2':(6,12,18),'azpeak':(13,30,45),'azgrow':(.5,1.5,2.5),'ppc':(.2,.15,.12),'sroas':(2,3,4),'shock':(1.05,1,1)}
FIX=dict(start=50,videos=8,cap=400,frac1=11/30,tt_fee=.08,spark=.10,spark_m=4,cr_pct=.30,tt_ret=.03,sh_pct=.029,sh_fix=.30,sh_ret=.02,amz_ref=.15,amz_start=5,az0=5,azramp=6,
 nsku=4,ful=2.5,lag=.5,cover=1.3,moq=150,soft=(500,1000,1500),ins1=150,ins2=500,acct=300,adpct=0.0,tax=.25,t1=20000,t2=5000,launch=2249,tm=350)
def params(s=2,**over):
    P=dict(FIX); P.update({k:v[s-1] for k,v in SCEN.items()}); P.update(over); return P
def tier(sku,q,shock):
    p=PROD[sku]['p']; i=bisect.bisect_right(TIERS,q)-1
    return (p[i] if i>=0 else p[0])*shock
def run(P,NM=36):
    fp_price=sum(o[0]*o[2] for o in OFF); fp_ship=sum(o[1]*o[2] for o in OFF)
    upf={s:sum(o[2]*o[3][i] for o in OFF) for i,s in enumerate(SK)}
    sub_price=sum(o[0]*o[2] for o in SUBS); sub_ship=sum(o[1]*o[2] for o in SUBS)
    ups={s:sum(o[2]*o[3][i] for o in SUBS) for i,s in enumerate(SK)}
    az_price=sum(PROD[s]['ret']*PROD[s]['mix'] for s in SK); az_fba=sum(PROD[s]['fba']*PROD[s]['mix'] for s in SK)
    # demand pass (independent of stock)
    D=[]
    cr=0; subs=0
    for m in range(0,NM+2):
        cal=(10+m-1)%12+1  # month 0 = Oct
        yr=0 if m==0 else (m-1)//12+1
        frac=0 if m==0 else (P['frac1'] if m==1 else 1)
        seas=SEAS[cal-1]
        if m==0: cr=0
        elif m<=2: cr=P['start']
        elif m<=12: cr=cr+P['add1']
        else: cr=min(P['cap'],cr+P['add2'])
        vid=cr*P['videos']*P['deliv']*frac; bl=(min(cr,50)*P['v1']+max(0,min(cr,200)-50)*P['v2']+max(0,cr-200)*P['v3'])/max(cr,1); views=vid*bl
        tt=views*P['conv']*(1+P['org'])*seas*(1+(P.get('spark',0) if m>=P.get('spark_m',3) else 0)*P.get('sroas',3))
        if m in P.get('tt_zero',()): tt=0
        sd=tt*P['direct']; new=(tt+sd)*P['subconv']
        subord=0 if m==0 else subs
        subs=0 if m==0 else subs*(1-P['churn'])+new
        k=m-P['amz_start']
        azday=0 if m<P['amz_start'] else (P['az0']+(P['azpeak']-P['az0'])*k/P['azramp'] if k<=P['azramp'] else P['azpeak']+P['azgrow']*(k-P['azramp']))
        az=azday*P['nsku']*30.4
        if m in P.get('all_zero',()): tt=sd=new=0; subord=0; az=0
        units={s:(tt+sd)*upf[s]+subord*ups[s]+az*PROD[s]['mix'] for s in SK}
        D.append(dict(m=m,yr=yr,cr=cr,vid=vid,views=views,tt=tt,sd=sd,new=new,subs=subs,subord=subord,az=az,units=units))
    out=[]; stock={s:0 for s in SK}; cash=0; prev=None
    for m in range(0,NM+1):
        d=D[m]; nx=D[m+1] if m<NM else D[m]
        if m==NM: nx=D[m]
        cost={s:tier(s,d['units'][s],P['shock']) for s in SK}
        r_tt=d['tt']*fp_price; r_sf=d['sd']*fp_price; r_ss=d['subord']*sub_price; r_az=d['az']*az_price; rev=r_tt+r_sf+r_ss+r_az
        cogs=sum(d['units'][s]*cost[s] for s in SK)
        ttfee=P['tt_fee']*r_tt; crc=P['cr_pct']*r_tt; shfee=P['sh_pct']*(r_sf+r_ss)+P['sh_fix']*(d['sd']+d['subord'])
        azfee=P['amz_ref']*r_az+d['az']*az_fba; ppc=P['ppc']*r_az
        ship=(d['tt']+d['sd'])*fp_ship+d['subord']*sub_ship; fulf=P['ful']*(d['tt']+d['sd']+d['subord'])
        ret=P['tt_ret']*r_tt+P['sh_ret']*(r_sf+r_ss)
        var=cogs+ttfee+crc+shfee+azfee+ppc+ship+fulf+ret; contrib=rev-var
        soft=0 if d['yr']==0 else P['soft'][d['yr']-1]
        ins=0 if m==0 else (P['ins1'] if m<12 else P['ins2']); acct=0 if m==0 else P['acct']
        team=0 if m==0 else sum(c for c,s0 in P.get('team',TEAM) if s0<=m)
        sp=(P.get('spark',0) if m>=P.get('spark_m',3) else 0); ads=sp*r_tt/(1+sp*P.get('sroas',3))+P['adpct']*rev; opex=soft+ins+acct+team+ads; ebitda=contrib-opex; tax=max(0,ebitda)*P['tax']; net=ebitda-tax
        launch=(P['launch']+P['tm']) if m==0 else 0
        tt_net=r_tt-ttfee-crc-P['tt_ret']*r_tt; az_net=r_az-azfee-ppc
        if m==0: in_tt=0; in_az=0
        else: in_tt=tt_net*(1-P['lag'])+prev['tt_net']*P['lag']; in_az=az_net*(1-P['lag'])+prev['az_net']*P['lag']
        in_sh=r_sf+r_ss-shfee-P['sh_ret']*(r_sf+r_ss)
        buy=0; pq={}
        for s in SK:
            op=stock[s]
            if m==0: q=P.get('fb_'+s,PROD[s]['fb'])
            else:
                need=P['cover']*nx['units'][s]; pre=op-d['units'][s]
                q=max(P['moq'],math.ceil((need-pre)/50)*50) if need>pre else 0
            pq[s]=q; buy+= q*tier(s,q,(1 if m==0 else P['shock'])) if q else 0
            stock[s]=op-(d['units'][s] if m>0 else 0)+q
        fund=P['t1'] if m==0 else (P['t2'] if m==1 else 0)
        ncf=in_tt+in_sh+in_az-(ship+fulf)-opex-tax-buy-launch+fund
        cash+=ncf
        row=dict(d); row.update(rev=rev,r_tt=r_tt,r_sf=r_sf,r_ss=r_ss,r_az=r_az,cogs=cogs,ttfee=ttfee,crc=crc,shfee=shfee,azfee=azfee,ppc=ppc,ship=ship,fulf=fulf,ret=ret,var=var,contrib=contrib,
          soft=soft,ins=ins,acct=acct,team=team,ads=ads,opex=opex,ebitda=ebitda,tax=tax,net=net,buy=buy,cash=cash,tt_net=tt_net,az_net=az_net,pq=pq,stock=dict(stock),cost=cost)
        out.append(row); prev=row
    return out
def year(o,y,k): return sum(r[k] for r in o if r['yr']==y)
def summary(o):
    S={}
    for y in (1,2,3):
        for k in ('rev','contrib','opex','ebitda','net','buy','tt','sd','subord','az','r_tt','r_sf','r_ss','r_az','cogs','crc','team'):
            S[f'{k}{y}']=year(o,y,k)
        S[f'subs{y}']=o[12*y]['subs']; S[f'cash{y}']=o[12*y]['cash']; S[f'cr{y}']=o[12*y]['cr']
    S['mincash']=min(r['cash'] for r in o); S['mincash_m']=min(range(len(o)),key=lambda i:o[i]['cash'])
    S['m1m']=next((r['m'] for r in o if r['m']>0 and r['ebitda']*12>=1e6),None)
    cum=0; S['payback']=None
    for r in o[1:]:
        cum+=r['ebitda']
        if cum>=FIX['t1']+FIX['t2']: S['payback']=r['m']; break
    return S
def verify(path='Aldera_FinModel.xlsx'):
    import json
    from openpyxl import load_workbook
    L=json.load(open('layout.json')); wb=load_workbook(path,data_only=True); M=wb['Модель']; rows=L['rows']
    o=run(params(2)); worst=0
    for k in ('rev','cogs','contrib','opex','ebitda','cash','buy' ):
        rk={'buy':'o_buy'}.get(k,k)
        for m in range(0,37):
            xv=M.cell(rows[rk],3+m).value or 0; pv=o[m][k]
            worst=max(worst,abs(xv-pv))
    return worst
if __name__=='__main__':
    print('max abs diff vs Excel:',round(verify(),6))
    S=summary(run(params(2))); print({k:round(v) if isinstance(v,float) else v for k,v in S.items() if k[-1] in '1' or k in ('mincash','m1m','payback','mincash_m')})
