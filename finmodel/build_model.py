"""Build The Aldera financial model (xlsx, live formulas)."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as CL
from openpyxl.comments import Comment

BLUE=Font(name='Arial',size=10,color='0000FF'); BLK=Font(name='Arial',size=10); BOLD=Font(name='Arial',size=10,bold=True)
GRN=Font(name='Arial',size=10,color='008000'); TITLE=Font(name='Arial',size=14,bold=True); H2=Font(name='Arial',size=11,bold=True,color='2C6A55')
YEL=PatternFill('solid',fgColor='FFFF00'); HDR=PatternFill('solid',fgColor='E2EEE8'); TOT=PatternFill('solid',fgColor='F1F5F2')
USD='$#,##0;($#,##0);"-"'; USD2='$#,##0.00;($#,##0.00);"-"'; PCT='0.0%;(0.0%);"-"'; NUM='#,##0;(#,##0);"-"'; NUM1='#,##0.0;(#,##0.0);"-"'
thin=Side(style='thin',color='D6DDD8')

wb=Workbook()
A=wb.active; A.title='Допущения'
REF={}
def ref(sheet,cell): return f"'{sheet}'!${''.join(c for c in cell if c.isalpha())}${''.join(c for c in cell if c.isdigit())}"
def setc(ws,cell,val,font=BLK,fmt=None,fill=None,note=None):
    c=ws[cell]; c.value=val; c.font=font
    if fmt: c.number_format=fmt
    if fill: c.fill=fill
    if note: c.comment=Comment(note,'Claude')
    return c

A['A1']='The Aldera · 518 GROUP LLC · Финансовая модель · допущения'; A['A1'].font=TITLE
A['A2']='Синий шрифт: вводные, их можно менять. Чёрный: формулы. Жёлтая заливка: ключевые рычаги. Модель пересчитывается сама.'; A['A2'].font=Font(name='Arial',size=9,italic=True)
setc(A,'A3','Сценарий: 1 консервативный, 2 базовый, 3 оптимистичный',BOLD)
setc(A,'B3',2,BLUE,fill=YEL,note='Переключатель сценария. Меняет все параметры в таблице ниже.')
REF['scen']=ref('Допущения','B3')

r=5
for i,h in enumerate(['Параметр сценария','Консерв.','Базовый','Оптим.','Активное','Пояснение']):
    c=A.cell(r,1+i,h); c.font=BOLD; c.fill=HDR
SCEN=[
 ('views','Просмотров на ролик',25000,25000,50000,NUM,'Решение владельца 03.10: минимум 25 000 просмотров на ролик'),
 ('conv','Конверсия просмотра в заказ',0.0003,0.0005,0.0007,'0.00%','Доля просмотров, которая становится заказом'),
 ('deliv','Доля роликов, которые реально выходят',0.60,1.00,1.00,PCT,'8 роликов по договору; база по условию владельца'),
 ('org','Органика и поиск сверху',0.10,0.20,0.30,PCT,'Заказы без ролика: поиск, профиль, повторный показ'),
 ('direct','Прямые заказы на Shopify, доля от TikTok',0.05,0.10,0.15,PCT,'Покупатели из TikTok, которые идут на сайт'),
 ('subconv','Переход первых покупателей в подписку',0.12,0.15,0.20,PCT,'Главный рычаг модели'),
 ('churn','Отток подписки в месяц',0.12,0.10,0.08,PCT,'Обычная цифра для подписки на добавки около 10%'),
 ('add1','Новых криейторов в месяц, мес. 3-12',5,10,15,NUM,'Старт 50 криейторов'),
 ('add2','Новых криейторов в месяц, с 13-го',6,12,18,NUM,''),
 ('azpeak','Amazon: продаж в день на продукт после разгона',13,30,45,NUM1,'Достигается через N месяцев разгона'),
 ('azgrow','Amazon: прирост в день на продукт в месяц после разгона',0.5,1.5,2.5,NUM1,''),
 ('ppc','Amazon: реклама PPC, доля цены',0.20,0.15,0.12,PCT,'Половина продаж через рекламу при ACoS 24-40%'),
 ('sroas','Spark Ads: выручка на $1 рекламы (ROAS)',2,3,4,'0.0x','Реклама идёт только на ролики с доказанной конверсией'),
 ('shock','Множитель себестоимости Vox',1.05,1.00,1.00,'0.00x','Рост цен поставщика'),
]
r=6
for key,lab,c1,c2,c3,fmt,note in SCEN:
    setc(A,f'A{r}',lab)
    for col,v in zip('BCD',(c1,c2,c3)): setc(A,f'{col}{r}',v,BLUE,fmt)
    setc(A,f'E{r}',f'=CHOOSE($B$3,B{r},C{r},D{r})',BLK,fmt,fill=YEL)
    setc(A,f'F{r}',note,Font(name='Arial',size=9,color='5A6762'))
    REF[key]=ref('Допущения',f'E{r}'); r+=1

r+=1; setc(A,f'A{r}','Постоянные допущения',H2); r+=1
FIX=[
 ('start','Криейторов на старте',50,NUM,'Решение владельца'),
 ('videos','Роликов на криейтора в месяц по договору',8,NUM,'Решение владельца: 8 роликов за 30%'),
 ('cap','Потолок числа криейторов',400,NUM,'Предел для одного менеджера-отдела'),
 ('frac1','Доля первого месяца (продажи с 20.11)',11/30,PCT,'11 дней из 30'),
 ('tt_fee','TikTok Shop: комиссия площадки',0.08,PCT,'С 04.08.2026 TikTok поднял ставку с 6% до 8% для большинства непищевых категорий. Ставку для добавок проверить в Seller Center'),
 ('spark','Spark Ads на выстрелившие ролики, доля органической выручки TikTok',0.10,PCT,'Решение владельца 03.10: платная реклама только на ролики, которые уже выстрелили'),
 ('spark_m','Spark Ads: с месяца модели',4,NUM,'Февраль 2027: после декабрьского провала денег'),
 ('cr_pct','Комиссия криейтора',0.30,PCT,'Решение владельца'),
 ('tt_ret','TikTok: возвраты, доля выручки',0.03,PCT,''),
 ('sh_pct','Shopify Payments: процент',0.029,PCT,'Тариф Shopify Basic'),
 ('sh_fix','Shopify Payments: за транзакцию',0.30,USD2,''),
 ('sh_ret','Shopify: возвраты, доля выручки',0.02,PCT,''),
 ('amz_ref','Amazon: комиссия за продажу',0.15,PCT,'Категория Health & Household'),
 ('amz_start','Amazon: месяц старта (1 = ноябрь 2026)',5,NUM,'Март 2027, после Brand Registry и допуска'),
 ('az0','Amazon: продаж в день на продукт в первый месяц',5,NUM1,''),
 ('azramp','Amazon: месяцев разгона',6,NUM,''),
 ('nsku','Число продуктов',4,NUM,''),
 ('ful','Склад-отправщик, за заказ TikTok и Shopify',2.50,USD2,'Пик-пак 3PL во Флориде, оценка'),
 ('lag','Доля выплат TikTok и Amazon, приходящая в следующем месяце',0.50,PCT,'Выплаты после доставки и раз в две недели'),
 ('cover','Запас товара после перезаказа, месяцев продаж',1.3,'0.0','Политика перезаказа'),
 ('moq','Минимальный заказ Vox, банок',150,NUM,'Прайс Vox 2026 Q3'),
 ('soft1','Сервисы и ПО в месяц, год 1',500,USD,'Shopify, подписка, почта, дизайн-сервисы'),
 ('soft2','Сервисы и ПО в месяц, год 2',1000,USD,''),
 ('soft3','Сервисы и ПО в месяц, год 3',1500,USD,''),
 ('ins1','Страховка ответственности в месяц, первые 11 мес.',150,USD,'Product liability'),
 ('ins2','Страховка ответственности в месяц, с 12-го',500,USD,'Amazon требует при продажах от $10 тыс. в месяц'),
 ('acct','Бухгалтерия и юристы в месяц',300,USD,''),
 ('adpct','Прочая реклама (Meta, Google), доля выручки',0.0,PCT,'Не планируется; 0 = не включено'),
 ('tax','Налог на прибыль, оценка',0.25,PCT,'Оценка: федеральный и штатный налог на прибыль LLC через владельца. Уточнить у бухгалтера'),
 ('t1','Вложение владельца, транш 1 (октябрь 2026)',20000,USD,'Решение владельца 03.10'),
 ('t2','Вложение владельца, транш 2 (конец ноября 2026)',5000,USD,'Решение владельца 03.10: партиями'),
 ('launch','Запуск: дизайн 4 этикеток, GS1, фрахт, образцы',2249,USD,'$499 + $300 + $500 + $950'),
 ('tm','Заявка USPTO на знак THE ALDERA, класс 5',350,USD,'Пошлина USPTO за 1 класс'),
]
for key,lab,v,fmt,note in FIX:
    setc(A,f'A{r}',lab); setc(A,f'B{r}',v,BLUE,fmt); setc(A,f'C{r}',note,Font(name='Arial',size=9,color='5A6762'))
    REF[key]=ref('Допущения',f'B{r}'); r+=1

# products
r+=1; setc(A,f'A{r}','Продукты (Vox Nutrition, сток, минимум 150)',H2); r+=1
TIERS=[150,300,500,1000,2500,5000,7500,10000]
hdr=['Продукт','Артикул Vox','Розница','Доставка 1 банки','FBA за банку','Доля в Amazon','Первая закупка, банок']+[f'Цена @{t}' for t in TIERS]
for i,h in enumerate(hdr): c=A.cell(r,1+i,h); c.font=BOLD; c.fill=HDR
r+=1; setc(A,f'A{r}','Порог ступени, банок',Font(name='Arial',size=9,italic=True))
for i,t in enumerate(TIERS): setc(A,f'{CL(8+i)}{r}',t,BLUE,NUM)
thr_row=r; REF['thr']=f"'Допущения'!$H${r}:$O${r}"; r+=1
PROD=[('COL','Colostrum','14721',39.95,4.50,6.55,0.35,1000,[8.52,8.37,8.12,7.97,7.82,7.67,7.42,7.12]),
      ('PRO',"Women's Probiotic",'14648',34.95,3.50,5.40,0.15,400,[8.34,7.99,7.64,7.44,7.24,6.94,6.74,6.59]),
      ('MG','Magnesium Glycinate','14683',29.95,3.50,5.40,0.25,1000,[5.51,5.16,4.81,4.61,4.41,4.11,3.91,3.76]),
      ('CRE','Creatine','14589',29.95,5.50,6.55,0.25,200,[6.90,6.75,6.50,6.35,6.20,6.05,5.80,5.50])]
PROW={}
for k,n,art,ret,sh,fba,mix,fb,prices in PROD:
    setc(A,f'A{r}',n,BOLD); setc(A,f'B{r}',art,BLUE)
    for col,v,fmt in (('C',ret,USD2),('D',sh,USD2),('E',fba,USD2),('F',mix,PCT),('G',fb,NUM)): setc(A,f'{col}{r}',v,BLUE,fmt)
    for i,p in enumerate(prices): setc(A,f'{CL(8+i)}{r}',p,BLUE,USD2)
    PROW[k]=r; r+=1
A[f'A{r}']='Источник цен: прайс Vox Nutrition 2026 Q3 Pricing Sheet. Розница: решение по позиционированию 03.10. Креатин в первой закупке 200, чтобы транш $20 000 покрыл и заявку на знак.'; A[f'A{r}'].font=Font(name='Arial',size=9,italic=True); r+=2
SK=['COL','PRO','MG','CRE']

# offers
setc(A,f'A{r}','Первая покупка: предложения и доли заказов',H2); r+=1
for i,h in enumerate(['Предложение','Цена','Доставка заказа','Доля заказов','Colostrum, банок','Probiotic, банок','Magnesium, банок','Creatine, банок']):
    c=A.cell(r,1+i,h); c.font=BOLD; c.fill=HDR
r+=1; off0=r
OFF=[('Foundation: Colostrum + Magnesium',59.95,5.50,.35,(1,0,1,0)),('Gut Duo: Colostrum + Probiotic',64.95,5.50,.10,(1,1,0,0)),
     ('The Daily Four',114.95,8.50,.05,(1,1,1,1)),('Colostrum',39.95,4.50,.20,(1,0,0,0)),("Women's Probiotic",34.95,3.50,.12,(0,1,0,0)),
     ('Magnesium Glycinate',29.95,3.50,.10,(0,0,1,0)),('Creatine',29.95,5.50,.08,(0,0,0,1))]
for n,p,sh,w,q in OFF:
    setc(A,f'A{r}',n); setc(A,f'B{r}',p,BLUE,USD2); setc(A,f'C{r}',sh,BLUE,USD2); setc(A,f'D{r}',w,BLUE,PCT)
    for i,x in enumerate(q): setc(A,f'{CL(5+i)}{r}',x,BLUE,NUM)
    r+=1
off1=r-1
setc(A,f'A{r}','Средний первый заказ',BOLD); setc(A,f'B{r}',f'=SUMPRODUCT(B{off0}:B{off1},$D${off0}:$D${off1})',BOLD,USD2)
setc(A,f'C{r}',f'=SUMPRODUCT(C{off0}:C{off1},$D${off0}:$D${off1})',BOLD,USD2); setc(A,f'D{r}',f'=SUM(D{off0}:D{off1})',BOLD,PCT)
for i in range(4): setc(A,f'{CL(5+i)}{r}',f'=SUMPRODUCT({CL(5+i)}{off0}:{CL(5+i)}{off1},$D${off0}:$D${off1})',BOLD,'0.000')
REF['fp_price']=ref('Допущения',f'B{r}'); REF['fp_ship']=ref('Допущения',f'C{r}')
for i,s in enumerate(SK): REF['upf_'+s]=ref('Допущения',f'{CL(5+i)}{r}')
A[f'I{r}']='← банок каждого продукта на один средний заказ'; A[f'I{r}'].font=Font(name='Arial',size=9,italic=True)
r+=2
setc(A,f'A{r}','Подписка Shopify, ежемесячно',H2); r+=1
for i,h in enumerate(['Подписка','Цена в месяц','Доставка','Доля подписчиков','Colostrum, банок','Probiotic, банок','Magnesium, банок','Creatine, банок']):
    c=A.cell(r,1+i,h); c.font=BOLD; c.fill=HDR
r+=1; s0=r
SUBS=[('Foundation',54.95,5.50,.60,(1,0,1,0)),('Gut Duo',59.95,5.50,.25,(1,1,0,0)),('The Daily Four, креатин раз в 2 мес.',99.95,7.50,.15,(1,1,1,0.5))]
for n,p,sh,w,q in SUBS:
    setc(A,f'A{r}',n); setc(A,f'B{r}',p,BLUE,USD2); setc(A,f'C{r}',sh,BLUE,USD2); setc(A,f'D{r}',w,BLUE,PCT)
    for i,x in enumerate(q): setc(A,f'{CL(5+i)}{r}',x,BLUE,'0.0')
    r+=1
s1=r-1
setc(A,f'A{r}','Средний подписчик',BOLD); setc(A,f'B{r}',f'=SUMPRODUCT(B{s0}:B{s1},$D${s0}:$D${s1})',BOLD,USD2)
setc(A,f'C{r}',f'=SUMPRODUCT(C{s0}:C{s1},$D${s0}:$D${s1})',BOLD,USD2); setc(A,f'D{r}',f'=SUM(D{s0}:D{s1})',BOLD,PCT)
for i in range(4): setc(A,f'{CL(5+i)}{r}',f'=SUMPRODUCT({CL(5+i)}{s0}:{CL(5+i)}{s1},$D${s0}:$D${s1})',BOLD,'0.000')
REF['sub_price']=ref('Допущения',f'B{r}'); REF['sub_ship']=ref('Допущения',f'C{r}')
for i,s in enumerate(SK): REF['ups_'+s]=ref('Допущения',f'{CL(5+i)}{r}')
r+=2
setc(A,f'A{r}','Amazon: средняя цена заказа',BOLD); setc(A,f'B{r}',f'=SUMPRODUCT(C{PROW["COL"]}:C{PROW["CRE"]},F{PROW["COL"]}:F{PROW["CRE"]})',BOLD,USD2); REF['az_price']=ref('Допущения',f'B{r}'); r+=1
setc(A,f'A{r}','Amazon: средний FBA за заказ',BOLD); setc(A,f'B{r}',f'=SUMPRODUCT(E{PROW["COL"]}:E{PROW["CRE"]},F{PROW["COL"]}:F{PROW["CRE"]})',BOLD,USD2); REF['az_fba']=ref('Допущения',f'B{r}'); r+=2

setc(A,f'A{r}','Сезонность по календарным месяцам',H2); r+=1
MONTHS=['январь','февраль','март','апрель','май','июнь','июль','август','сентябрь','октябрь','ноябрь','декабрь']
SEAS=[1.10,1,1,1,1,1,1,1,1,1,1.15,1.15]
se0=r
for mname,f in zip(MONTHS,SEAS): setc(A,f'A{r}',mname); setc(A,f'B{r}',f,BLUE,'0.00'); r+=1
REF['seas']=f"'Допущения'!$B${se0}:$B${r-1}"
A[f'C{se0}']='Black Friday и праздники +15%, январь (новогодние цели) +10%'; A[f'C{se0}'].font=Font(name='Arial',size=9,italic=True)
r+=1
setc(A,f'A{r}','Команда (найм по месяцам модели)',H2); r+=1
for i,h in enumerate(['Роль','Стоимость в месяц','С месяца модели']): c=A.cell(r,1+i,h); c.font=BOLD; c.fill=HDR
r+=1; t0=r
TEAM=[('Менеджер криейторов 1',4000,5),('Поддержка клиентов',3000,6),('Менеджер Amazon и операций',5000,10),('Менеджер криейторов 2',4000,13),('Поддержка клиентов 2',3000,16),('Финансы, на аутсорсе',2000,18)]
for n,c,m in TEAM: setc(A,f'A{r}',n); setc(A,f'B{r}',c,BLUE,USD); setc(A,f'C{r}',m,BLUE,NUM); r+=1
REF['team_cost']=f"'Допущения'!$B${t0}:$B${r-1}"; REF['team_start']=f"'Допущения'!$C${t0}:$C${r-1}"
A[f'D{t0}']='Месяц 1 = ноябрь 2026. Зарплата владельца не заложена.'; A[f'D{t0}'].font=Font(name='Arial',size=9,italic=True)
A.column_dimensions['A'].width=58; A.column_dimensions['F'].width=50
for col in 'BCDE': A.column_dimensions[col].width=14
for i in range(8): A.column_dimensions[CL(8+i)].width=11
A.column_dimensions['C'].width=16

# ---------------- MODEL ----------------
M=wb.create_sheet('Модель')
NM=36; C0=3  # month m in column C0+m
def col(m): return CL(C0+m)
M['A1']='Помесячная модель, 36 месяцев (ноябрь 2026 - октябрь 2029)'; M['A1'].font=TITLE
rows={}
def row(key,label,rownum,fmt=NUM,bold=False):
    rows[key]=rownum; c=M.cell(rownum,1,label); c.font=BOLD if bold else BLK
    for m in range(0,NM+1): M.cell(rownum,C0+m).number_format=fmt
    return rownum
def fill(key,f,mstart=0,font=BLK):
    rr=rows[key]
    for m in range(mstart,NM+1):
        cl=col(m); pl=col(m-1) if m>0 else None; nx=col(min(m+1,NM))
        v=f(m,cl,pl,nx)
        if v is None: continue
        c=M[f'{cl}{rr}']; c.value=v; c.font=BOLD if M.cell(rr,1).font.bold else font
R=lambda k: rows[k]
row('m','Месяц модели',2,'0'); fill('m',lambda m,c,p,n: 0 if m==0 else f'={p}2+1')
row('date','Месяц',3,'mmm yy'); fill('date',lambda m,c,p,n: f'=DATE(2026,10+{c}2,1)')
row('cal','Календарный месяц',4,'0'); fill('cal',lambda m,c,p,n: f'=MONTH({c}3)')
row('yr','Год модели',5,'0'); fill('yr',lambda m,c,p,n: f'=IF({c}2=0,0,INT(({c}2-1)/12)+1)')
row('frac','Доля месяца',6,'0.00'); fill('frac',lambda m,c,p,n: f'=IF({c}2=0,0,IF({c}2=1,{REF["frac1"]},1))')
row('seas','Сезонность',7,'0.00'); fill('seas',lambda m,c,p,n: f'=INDEX({REF["seas"]},{c}4)')
M['A8']='Воронка'; M['A8'].font=H2
row('cr','Активных криейторов',9,NUM,True); fill('cr',lambda m,c,p,n: 0 if m==0 else f'=IF({c}2<=2,{REF["start"]},IF({c}2<=12,{p}9+{REF["add1"]},MIN({REF["cap"]},{p}9+{REF["add2"]})))')
row('vid','Роликов вышло',10); fill('vid',lambda m,c,p,n: f'={c}9*{REF["videos"]}*{REF["deliv"]}*{c}6')
row('views','Просмотров',11); fill('views',lambda m,c,p,n: f'={c}10*{REF["views"]}')
row('tt','Заказы TikTok (первые покупки)',12,NUM,True); fill('tt',lambda m,c,p,n: f'={c}11*{REF["conv"]}*(1+{REF["org"]})*{c}7*(1+IF({c}2>={REF["spark_m"]},{REF["spark"]},0)*{REF["sroas"]})')
row('sd','Заказы Shopify прямые (первые)',13); fill('sd',lambda m,c,p,n: f'={c}12*{REF["direct"]}')
row('newsub','Новые подписчики',14); fill('newsub',lambda m,c,p,n: f'=({c}12+{c}13)*{REF["subconv"]}')
row('subs','Подписчиков на конец месяца',15,NUM,True); fill('subs',lambda m,c,p,n: 0 if m==0 else f'={p}15*(1-{REF["churn"]})+{c}14')
row('subord','Заказы подписки (списание)',16); fill('subord',lambda m,c,p,n: 0 if m==0 else f'={p}15')
row('azday','Amazon: продаж в день на продукт',17,NUM1); fill('azday',lambda m,c,p,n: f'=IF({c}2<{REF["amz_start"]},0,IF({c}2-{REF["amz_start"]}<={REF["azramp"]},{REF["az0"]}+({REF["azpeak"]}-{REF["az0"]})*({c}2-{REF["amz_start"]})/{REF["azramp"]},{REF["azpeak"]}+{REF["azgrow"]}*({c}2-{REF["amz_start"]}-{REF["azramp"]})))')
row('az','Заказы Amazon',18,NUM,True); fill('az',lambda m,c,p,n: f'={c}17*{REF["nsku"]}*30.4')
row('ord','Всего заказов',19,NUM,True); fill('ord',lambda m,c,p,n: f'={c}12+{c}13+{c}16+{c}18')
M['A20']='Банок продано'; M['A20'].font=H2
NAMES={'COL':'Colostrum','PRO':"Women's Probiotic",'MG':'Magnesium Glycinate','CRE':'Creatine'}
for i,s in enumerate(SK):
    rr=21+i; row('u_'+s,f'Банок: {NAMES[s]}',rr)
    mixref=ref('Допущения',f'F{PROW[s]}')
    fill('u_'+s,lambda m,c,p,n,s=s,mixref=mixref: f'=({c}12+{c}13)*{REF["upf_"+s]}+{c}16*{REF["ups_"+s]}+{c}18*{mixref}')
row('units','Банок всего',25,NUM,True); fill('units',lambda m,c,p,n: f'=SUM({c}21:{c}24)')
M['A26']='Себестоимость банки (ступень по объёму месяца)'; M['A26'].font=H2
for i,s in enumerate(SK):
    rr=27+i; row('c_'+s,f'Банка: {NAMES[s]}',rr,USD2); pr=PROW[s]
    fill('c_'+s,lambda m,c,p,n,pr=pr,rr=rr,i=i: f'=IFERROR(INDEX(\'Допущения\'!$H${pr}:$O${pr},MATCH({c}{21+i},{REF["thr"]},1)),\'Допущения\'!$H${pr})*{REF["shock"]}')
M['A31']='Выручка'; M['A31'].font=H2
row('r_tt','TikTok Shop',32,USD); fill('r_tt',lambda m,c,p,n: f'={c}12*{REF["fp_price"]}')
row('r_sf','Shopify: первые покупки',33,USD); fill('r_sf',lambda m,c,p,n: f'={c}13*{REF["fp_price"]}')
row('r_ss','Shopify: подписка',34,USD); fill('r_ss',lambda m,c,p,n: f'={c}16*{REF["sub_price"]}')
row('r_az','Amazon',35,USD); fill('r_az',lambda m,c,p,n: f'={c}18*{REF["az_price"]}')
row('rev','Выручка всего',36,USD,True); fill('rev',lambda m,c,p,n: f'=SUM({c}32:{c}35)')
M['A37']='Переменные расходы'; M['A37'].font=H2
row('cogs','Товар (Vox)',38,USD); fill('cogs',lambda m,c,p,n: f'={c}21*{c}27+{c}22*{c}28+{c}23*{c}29+{c}24*{c}30')
row('ttfee','Комиссия TikTok Shop',39,USD); fill('ttfee',lambda m,c,p,n: f'={REF["tt_fee"]}*{c}32')
row('crc','Комиссия криейторов',40,USD); fill('crc',lambda m,c,p,n: f'={REF["cr_pct"]}*{c}32')
row('shfee','Эквайринг Shopify',41,USD); fill('shfee',lambda m,c,p,n: f'={REF["sh_pct"]}*({c}33+{c}34)+{REF["sh_fix"]}*({c}13+{c}16)')
row('azfee','Amazon: комиссия и FBA',42,USD); fill('azfee',lambda m,c,p,n: f'={REF["amz_ref"]}*{c}35+{c}18*{REF["az_fba"]}')
row('ppc','Amazon: реклама PPC',43,USD); fill('ppc',lambda m,c,p,n: f'={REF["ppc"]}*{c}35')
row('ship','Доставка покупателю',44,USD); fill('ship',lambda m,c,p,n: f'=({c}12+{c}13)*{REF["fp_ship"]}+{c}16*{REF["sub_ship"]}')
row('fulf','Склад-отправщик',45,USD); fill('fulf',lambda m,c,p,n: f'={REF["ful"]}*({c}12+{c}13+{c}16)')
row('ret','Возвраты',46,USD); fill('ret',lambda m,c,p,n: f'={REF["tt_ret"]}*{c}32+{REF["sh_ret"]}*({c}33+{c}34)')
row('var','Переменные расходы всего',47,USD,True); fill('var',lambda m,c,p,n: f'=SUM({c}38:{c}46)')
row('contrib','Маржинальная прибыль',48,USD,True); fill('contrib',lambda m,c,p,n: f'={c}36-{c}47')
row('cm','Маржинальность',49,PCT); fill('cm',lambda m,c,p,n: f'=IF({c}36=0,0,{c}48/{c}36)')
M['A50']='Постоянные расходы'; M['A50'].font=H2
row('soft','Сервисы и ПО',51,USD); fill('soft',lambda m,c,p,n: f'=IF({c}5=0,0,CHOOSE({c}5,{REF["soft1"]},{REF["soft2"]},{REF["soft3"]}))')
row('ins','Страховка',52,USD); fill('ins',lambda m,c,p,n: f'=IF({c}2=0,0,IF({c}2<12,{REF["ins1"]},{REF["ins2"]}))')
row('acct','Бухгалтерия и юристы',53,USD); fill('acct',lambda m,c,p,n: f'=IF({c}2=0,0,{REF["acct"]})')
row('team','Команда',54,USD); fill('team',lambda m,c,p,n: f'=IF({c}2=0,0,SUMPRODUCT(({REF["team_start"]}<={c}2)*{REF["team_cost"]}))')
row('ads','Реклама: Spark Ads на выстрелившие ролики и прочая',55,USD); fill('ads',lambda m,c,p,n: f'=IF({c}2>={REF["spark_m"]},{REF["spark"]},0)*{c}32/(1+IF({c}2>={REF["spark_m"]},{REF["spark"]},0)*{REF["sroas"]})+{REF["adpct"]}*{c}36')
row('opex','Постоянные расходы всего',56,USD,True); fill('opex',lambda m,c,p,n: f'=SUM({c}51:{c}55)')
row('ebitda','EBITDA (прибыль до налога)',57,USD,True); fill('ebitda',lambda m,c,p,n: f'={c}48-{c}56')
row('taxr','Налог на прибыль, оценка',58,USD); fill('taxr',lambda m,c,p,n: f'=MAX(0,{c}57)*{REF["tax"]}')
row('net','Чистая прибыль',59,USD,True); fill('net',lambda m,c,p,n: f'={c}57-{c}58')
row('launch','Запуск и знак (месяц 0)',60,USD); fill('launch',lambda m,c,p,n: f'=IF({c}2=0,{REF["launch"]}+{REF["tm"]},0)')
M['A61']='Денежный поток'; M['A61'].font=H2
row('in_tt','Поступления TikTok',62,USD); fill('in_tt',lambda m,c,p,n: 0 if m==0 else f'=({c}32-{c}39-{c}40-{REF["tt_ret"]}*{c}32)*(1-{REF["lag"]})+({p}32-{p}39-{p}40-{REF["tt_ret"]}*{p}32)*{REF["lag"]}')
row('in_sh','Поступления Shopify',63,USD); fill('in_sh',lambda m,c,p,n: f'={c}33+{c}34-{c}41-{REF["sh_ret"]}*({c}33+{c}34)')
row('in_az','Поступления Amazon',64,USD); fill('in_az',lambda m,c,p,n: 0 if m==0 else f'=({c}35-{c}42-{c}43)*(1-{REF["lag"]})+({p}35-{p}42-{p}43)*{REF["lag"]}')
row('cin','Поступления всего',65,USD,True); fill('cin',lambda m,c,p,n: f'=SUM({c}62:{c}64)')
row('o_ship','Доставка и склад',66,USD); fill('o_ship',lambda m,c,p,n: f'={c}44+{c}45')
row('o_opex','Постоянные расходы',67,USD); fill('o_opex',lambda m,c,p,n: f'={c}56')
row('o_tax','Налог',68,USD); fill('o_tax',lambda m,c,p,n: f'={c}58')
row('o_buy','Закупка товара у Vox',69,USD); fill('o_buy',lambda m,c,p,n: f'={c}77+{c}81+{c}85+{c}89')
row('o_launch','Запуск и знак',70,USD); fill('o_launch',lambda m,c,p,n: f'={c}60')
row('fund','Вложения владельца',71,USD); fill('fund',lambda m,c,p,n: f'=IF({c}2=0,{REF["t1"]},IF({c}2=1,{REF["t2"]},0))')
row('ncf','Чистый денежный поток',72,USD,True); fill('ncf',lambda m,c,p,n: f'={c}65-{c}66-{c}67-{c}68-{c}69-{c}70+{c}71')
row('cash','Деньги на счёте на конец месяца',73,USD,True); fill('cash',lambda m,c,p,n: f'={c}72' if m==0 else f'={p}73+{c}72')
M['A74']='Склад и перезаказы (заказ в месяце приходит к следующему)'; M['A74'].font=H2
for i,s in enumerate(SK):
    b=75+i*4; pr=PROW[s]; ur=21+i
    row('op_'+s,f'{NAMES[s]}: остаток на начало',b); fill('op_'+s,lambda m,c,p,n,b=b: 0 if m==0 else f'={p}{b+3}')
    row('pq_'+s,f'{NAMES[s]}: перезаказ, банок',b+1)
    fill('pq_'+s,lambda m,c,p,n,b=b,ur=ur,pr=pr: f"='Допущения'!$G${pr}" if m==0 else f'=IF({REF["cover"]}*{n}{ur}>{c}{b}-{c}{ur},MAX({REF["moq"]},CEILING({REF["cover"]}*{n}{ur}-({c}{b}-{c}{ur}),50)),0)')
    row('pc_'+s,f'{NAMES[s]}: перезаказ, $',b+2,USD)
    fill('pc_'+s,lambda m,c,p,n,b=b,pr=pr: f'=IF({c}{b+1}=0,0,{c}{b+1}*IFERROR(INDEX(\'Допущения\'!$H${pr}:$O${pr},MATCH({c}{b+1},{REF["thr"]},1)),\'Допущения\'!$H${pr})*IF({c}2=0,1,{REF["shock"]}))')
    row('cl_'+s,f'{NAMES[s]}: остаток на конец',b+3); fill('cl_'+s,lambda m,c,p,n,b=b,ur=ur: f'={c}{b}-{c}{ur}+{c}{b+1}')
M.column_dimensions['A'].width=46; M.column_dimensions['B'].width=4
for m in range(NM+1): M.column_dimensions[col(m)].width=11.5
M.freeze_panes='C4'
for cc in range(1,C0+NM+1):
    for rr in (2,3): M.cell(rr,cc).font=BOLD; M.cell(rr,cc).fill=HDR
M['B2']=''; 

# ------------- YEARS -------------
Y=wb.create_sheet('Итоги по годам',0)
Y['A1']='The Aldera · итоги по годам модели'; Y['A1'].font=TITLE
Y['A2']='Год 1: ноябрь 2026 - октябрь 2027. Год 2: ноябрь 2027 - октябрь 2028. Год 3: ноябрь 2028 - октябрь 2029. Сценарий на листе «Допущения», ячейка B3.'; Y['A2'].font=Font(name='Arial',size=9,italic=True)
for i,h in enumerate(['Показатель','Год 1','Год 2','Год 3','Итого 3 года']): c=Y.cell(4,1+i,h); c.font=BOLD; c.fill=HDR
rng=lambda r: f"'Модель'!$D${r}:${col(NM)}${r}"
yrng="'Модель'!$D$5:$"+col(NM)+"$5"
YL=[('Заказы TikTok','tt',NUM),('Заказы Shopify прямые','sd',NUM),('Заказы подписки','subord',NUM),('Заказы Amazon','az',NUM),('Всего заказов','ord',NUM),
    ('Банок продано','units',NUM),None,('Выручка TikTok Shop','r_tt',USD),('Выручка Shopify: первые','r_sf',USD),('Выручка Shopify: подписка','r_ss',USD),('Выручка Amazon','r_az',USD),('Выручка всего','rev',USD),None,
    ('Товар (Vox)','cogs',USD),('Комиссия TikTok Shop','ttfee',USD),('Комиссия криейторов','crc',USD),('Эквайринг Shopify','shfee',USD),('Amazon: комиссия и FBA','azfee',USD),('Amazon: реклама PPC','ppc',USD),
    ('Доставка покупателю','ship',USD),('Склад-отправщик','fulf',USD),('Возвраты','ret',USD),('Маржинальная прибыль','contrib',USD),None,
    ('Сервисы и ПО','soft',USD),('Страховка','ins',USD),('Бухгалтерия и юристы','acct',USD),('Команда','team',USD),('Реклама: Spark Ads и прочая','ads',USD),('Постоянные расходы','opex',USD),
    ('EBITDA (прибыль до налога)','ebitda',USD),('Налог, оценка','taxr',USD),('Чистая прибыль','net',USD),None,('Закупка товара у Vox','o_buy',USD)]
r=5; YROW={}
for it in YL:
    if it is None: r+=1; continue
    lab,key,fmt=it; Y.cell(r,1,lab).font=BLK
    for j in range(3):
        c=Y.cell(r,2+j,f'=SUMIF({yrng},{j+1},{rng(rows[key])})'); c.number_format=fmt; c.font=GRN
    c=Y.cell(r,5,f'=SUM(B{r}:D{r})'); c.number_format=fmt
    if key in ('rev','contrib','ebitda','net','ord'):
        for cc in range(1,6): Y.cell(r,cc).font=BOLD; Y.cell(r,cc).fill=TOT
    YROW[key]=r; r+=1
r+=1
Y.cell(r,1,'Маржинальность').font=BLK
for j in range(4): c=Y.cell(r,2+j,f'=IF({CL(2+j)}{YROW["rev"]}=0,0,{CL(2+j)}{YROW["contrib"]}/{CL(2+j)}{YROW["rev"]})'); c.number_format=PCT
YROW['cmp']=r; r+=1
Y.cell(r,1,'Рентабельность по EBITDA').font=BLK
for j in range(4): c=Y.cell(r,2+j,f'=IF({CL(2+j)}{YROW["rev"]}=0,0,{CL(2+j)}{YROW["ebitda"]}/{CL(2+j)}{YROW["rev"]})'); c.number_format=PCT
YROW['em']=r; r+=1
for lab,key,fmt in (('Подписчиков на конец года','subs',NUM),('Криейторов на конец года','cr',NUM),('Деньги на счёте на конец года','cash',USD)):
    Y.cell(r,1,lab)
    for j in range(3): c=Y.cell(r,2+j,f"=INDEX({rng(rows[key])},{12*(j+1)})"); c.number_format=fmt; c.font=GRN
    YROW[key+'_end']=r; r+=1
Y.cell(r,1,'Минимум денег на счёте за 3 года'); c=Y.cell(r,5,f"=MIN('Модель'!$C$73:${col(NM)}$73)"); c.number_format=USD; c.font=GRN; YROW['mincash']=r; r+=1
Y.cell(r,1,'Месяц выхода на темп $1 млн EBITDA в год'); c=Y.cell(r,5,f"=IFERROR(MATCH(TRUE,INDEX('Модель'!$D$57:${col(NM)}$57*12>=1000000,0),0),\"нет\")"); c.font=GRN; YROW['m1m']=r; r+=1
Y.cell(r,1,'Окупаемость вложений, месяц (накопленная EBITDA > вложений)')
Y.cell(r,5,'см. лист «Сводка»'); r+=1
Y.column_dimensions['A'].width=52
for cc in 'BCDE': Y.column_dimensions[cc].width=16
Y.freeze_panes='B5'

# ------------- UNIT ECONOMICS -------------
U=wb.create_sheet('Юнит-экономика',1)
U['A1']='Юнит-экономика: прибыль с банки и с заказа по каналам'; U['A1'].font=TITLE
hdr=['Продукт','Розница','Банка @1 000','Банка @5 000','Прибыль TikTok','Маржа TikTok','Прибыль Shopify','Маржа Shopify','Прибыль Amazon','Маржа Amazon']
for i,h in enumerate(hdr): c=U.cell(3,1+i,h); c.font=BOLD; c.fill=HDR
r=4
for s in SK:
    pr=PROW[s]; D=lambda cc: f"'Допущения'!${cc}${pr}"
    U.cell(r,1,NAMES[s]).font=BOLD
    U.cell(r,2,f"={D('C')}").number_format=USD2
    U.cell(r,3,f"={D('K')}*{REF['shock']}").number_format=USD2
    U.cell(r,4,f"={D('M')}*{REF['shock']}").number_format=USD2
    U.cell(r,5,f"=B{r}-C{r}-B{r}*({REF['tt_fee']}+{REF['cr_pct']}+{REF['tt_ret']})-{D('D')}-{REF['ful']}").number_format=USD2
    U.cell(r,6,f"=E{r}/B{r}").number_format=PCT
    U.cell(r,7,f"=B{r}-C{r}-B{r}*({REF['sh_pct']}+{REF['sh_ret']})-{REF['sh_fix']}-{D('D')}-{REF['ful']}").number_format=USD2
    U.cell(r,8,f"=G{r}/B{r}").number_format=PCT
    U.cell(r,9,f"=B{r}-C{r}-B{r}*({REF['amz_ref']}+{REF['ppc']})-{D('E')}").number_format=USD2
    U.cell(r,10,f"=I{r}/B{r}").number_format=PCT
    r+=1
r+=1
U.cell(r,1,'Первая покупка').font=H2; r+=1
for i,h in enumerate(['Предложение','Цена','Товар @1 000','Прибыль TikTok','Прибыль Shopify','Доля заказов']): c=U.cell(r,1+i,h); c.font=BOLD; c.fill=HDR
r+=1
costvec=lambda rowA: '+'.join(f"'Допущения'!{CL(5+i)}{rowA}*'Допущения'!$K${PROW[s]}*{REF['shock']}" for i,s in enumerate(SK))
for k in range(off0,off1+1):
    U.cell(r,1,f"='Допущения'!A{k}")
    U.cell(r,2,f"='Допущения'!B{k}").number_format=USD2
    U.cell(r,3,f"={costvec(k)}").number_format=USD2
    U.cell(r,4,f"=B{r}-C{r}-B{r}*({REF['tt_fee']}+{REF['cr_pct']}+{REF['tt_ret']})-'Допущения'!C{k}-{REF['ful']}").number_format=USD2
    U.cell(r,5,f"=B{r}-C{r}-B{r}*({REF['sh_pct']}+{REF['sh_ret']})-{REF['sh_fix']}-'Допущения'!C{k}-{REF['ful']}").number_format=USD2
    U.cell(r,6,f"='Допущения'!D{k}").number_format=PCT
    r+=1
fo0=r-(off1-off0+1); fo1=r-1
U.cell(r,1,'Средний первый заказ').font=BOLD
U.cell(r,2,f'=SUMPRODUCT(B{fo0}:B{fo1},F{fo0}:F{fo1})').number_format=USD2
U.cell(r,4,f'=SUMPRODUCT(D{fo0}:D{fo1},F{fo0}:F{fo1})').number_format=USD2
U.cell(r,5,f'=SUMPRODUCT(E{fo0}:E{fo1},F{fo0}:F{fo1})').number_format=USD2
avg_first_row=r; r+=2
U.cell(r,1,'Подписка Shopify').font=H2; r+=1
for i,h in enumerate(['Подписка','Цена в месяц','Товар @1 000','Прибыль в месяц','Прибыль в год','Доля','Жизнь подписчика, мес.','Прибыль за жизнь (LTV)']): c=U.cell(r,1+i,h); c.font=BOLD; c.fill=HDR
r+=1; us0=r
for k in range(s0,s1+1):
    U.cell(r,1,f"='Допущения'!A{k}")
    U.cell(r,2,f"='Допущения'!B{k}").number_format=USD2
    U.cell(r,3,f"={costvec(k)}").number_format=USD2
    U.cell(r,4,f"=B{r}-C{r}-B{r}*({REF['sh_pct']}+{REF['sh_ret']})-{REF['sh_fix']}-'Допущения'!C{k}-{REF['ful']}").number_format=USD2
    U.cell(r,5,f"=D{r}*12").number_format=USD
    U.cell(r,6,f"='Допущения'!D{k}").number_format=PCT
    U.cell(r,7,f"=1/{REF['churn']}").number_format='0.0'
    U.cell(r,8,f"=D{r}*G{r}").number_format=USD
    r+=1
U.cell(r,1,'Средний подписчик').font=BOLD
for cc,fmt in (('B',USD2),('D',USD2),('E',USD),('H',USD)):
    c=U[f'{cc}{r}']; c.value=f'=SUMPRODUCT({cc}{us0}:{cc}{r-1},F{us0}:F{r-1})'; c.number_format=fmt; c.font=BOLD
U[f'G{r}']=f'=1/{REF["churn"]}'; U[f'G{r}'].number_format='0.0'
avg_sub_row=r; r+=2
U.cell(r,1,'Ценность одного криейтора в месяц').font=H2; r+=1
lines=[('Заказов с роликов и органики',f"={REF['videos']}*{REF['deliv']}*{REF['views']}*{REF['conv']}*(1+{REF['org']})",NUM1),
       ('Продаж на сумму',f'=B{r}*\'Допущения\'!{REF["fp_price"].split("!")[1]}',USD),
       ('Заработок криейтора',f'=B{r+1}*{REF["cr_pct"]}',USD),
       ('Заработок криейтора за ролик',f'=B{r+2}/({REF["videos"]}*{REF["deliv"]})',USD2),
       ('Прибыль бренда с первых покупок',f'=B{r}*D{avg_first_row}',USD),
       ('Новых подписчиков',f'=B{r}*(1+{REF["direct"]})*{REF["subconv"]}',NUM1),
       ('Прибыль с этих подписчиков за их жизнь',f'=B{r+5}*H{avg_sub_row}',USD)]
for lab,f,fmt in lines:
    U.cell(r,1,lab); c=U.cell(r,2,f); c.number_format=fmt; r+=1
U.column_dimensions['A'].width=40
for cc in 'BCDEFGHIJ': U.column_dimensions[cc].width=15

wb.save('/home/user/The-Lucere/finmodel/Aldera_FinModel.xlsx')
import json; json.dump(dict(rows=rows,YROW=YROW,PROW=PROW,avg_first_row=avg_first_row,avg_sub_row=avg_sub_row),open('/home/user/The-Lucere/finmodel/layout.json','w'))
print('built')
